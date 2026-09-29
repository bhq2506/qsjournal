"""
Regenerate data/journals.json from the QS Subject / ASJC workbook.

Usage:
    pip install openpyxl
    python scripts/build_data.py path/to/QS_Subject_ASJCode_Journal_.xlsx

Expected workbook layout (as of Sep 2026):
  - 'SJR 2025'  : col A = Scopus Source ID, col B = SJR best quartile
  - 'Menu'      : hyperlinks to subject sheets (used for display names)
  - one sheet per QS subject: row1 = title, row2 = "QS Broad Faculty Area: ...",
    row3 = "ASJC Codes: ...", row6 = header, data from row 7
    (col A = Source ID, B = Title, C = ISSN, D = EISSN, J = OA status,
     S = Publisher, T = Publisher group)
"""
import json
import sys
from datetime import date

import openpyxl

SKIP = {"Menu", "QS Subject Areas", "All Journal", "SJR 2025"}
BROAD_ORDER = [
    "Arts and Humanities",
    "Engineering and Technology",
    "Life Sciences and Medicine",
    "Natural Sciences",
    "Social Sciences and Management",
]


def main(path, out="data/journals.json"):
    wb = openpyxl.load_workbook(path, read_only=True)

    sjr = {}
    for r in wb["SJR 2025"].iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            sjr[int(r[0])] = r[1]

    # Display names from the Menu sheet hyperlinks (needs non-read-only mode)
    wbl = openpyxl.load_workbook(path)
    # Menu columns B/D/F/H/J = the five QS broad faculty areas (official primary grouping)
    menu, menu_group = {}, {}
    col_group = {2: BROAD_ORDER[0], 4: BROAD_ORDER[1], 6: BROAD_ORDER[2], 8: BROAD_ORDER[3], 10: BROAD_ORDER[4]}
    for row in wbl["Menu"].iter_rows():
        for c in row:
            if c.hyperlink and c.hyperlink.location:
                sheet = c.hyperlink.location.split("!")[0].strip("'")
                menu[sheet] = c.value
                menu_group[sheet] = col_group.get(c.column)

    subjects, journals = [], {}
    for ws in wb.worksheets:
        if ws.title in SKIP:
            continue
        rows = ws.iter_rows(values_only=True)
        title = next(rows)[0]
        broad = next(rows)[0].split(":", 1)[1].strip()
        asjc = next(rows)[0].split(":", 1)[1].strip()
        for _ in range(3):
            next(rows)
        idx, n = len(subjects), 0
        for r in rows:
            if r[0] is None:
                continue
            sid = int(r[0])
            n += 1
            j = journals.get(sid)
            if not j:
                q = sjr.get(sid) or "Not ranked"
                if q in ("-", ""):
                    q = "Not ranked"
                j = journals[sid] = {
                    "id": sid, "t": r[1], "p": r[18] or r[19] or "", "q": q,
                    "issn": r[2] or "", "eissn": r[3] or "", "oa": r[9] or "", "s": [],
                }
            j["s"].append(idx)
        broads = [b.strip() for b in broad.split(";")]
        subjects.append({
            "name": menu.get(ws.title, title),
            "group": menu_group.get(ws.title) or broads[0],
            "broad": broads,
            "asjc": asjc, "n": n,
        })

    order = sorted(range(len(subjects)),
                   key=lambda i: (BROAD_ORDER.index(subjects[i]["group"]), subjects[i]["name"]))
    remap = {old: new for new, old in enumerate(order)}
    S = [subjects[i] for i in order]
    J = [[j["id"], j["t"], j["p"], j["q"], j["issn"], j["eissn"], j["oa"],
          sorted(remap[s] for s in j["s"])]
         for j in sorted(journals.values(), key=lambda j: j["t"].lower())]

    data = {
        "meta": {
            "source": "Scopus Source List + SJR 2025 best quartile; QS ASJC mapping",
            "generated": date.today().isoformat(),
            "journals": len(J),
        },
        "broad": BROAD_ORDER,
        "subjects": S,
        "journals": J,
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"Wrote {out}: {len(S)} subjects, {len(J)} journals")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1])
